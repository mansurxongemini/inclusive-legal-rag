"""
Local File Importer for Inclusive Legal RAG.
Automatically discovers, parses, vectorizes, and indexes Word (.doc, .docx), PDF (.pdf),
and text documents from data/qonunlar or custom directory.
"""
import hashlib
import logging
import os
import re
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from bs4 import BeautifulSoup
from pydantic import BaseModel, Field

from src.ingestion.localized import localized, provenance_for
from src.nlp.transliteration import cyrillic_to_latin, detect_script, normalize_apostrophes
from src.schema.canonical import (
    ArticleMetadata,
    ArticlePoint,
    ArticleStatus,
    CanonicalArticle,
    CrossReference,
    RefType,
    TranslationProvenanceType,
)
from src.storage.hybrid_store import HybridStore

logger = logging.getLogger(__name__)


class FileImportResult(BaseModel):
    filename: str
    file_type: str
    articles_count: int
    law_number: str
    law_title: str
    status: str
    error: Optional[str] = None


class BatchFileImportReport(BaseModel):
    success: bool
    total_files_scanned: int
    files_processed: int
    total_articles_indexed: int
    results: List[FileImportResult]
    message: str


class LocalFileImporter:
    """Discovers and parses local documents (.docx, .doc, .pdf, .txt) into canonical articles."""

    def __init__(self, store: Optional[HybridStore] = None):
        self.store = store or HybridStore()

    async def scan_and_import_directory(
        self,
        directory_path: str = "data/qonunlar",
        force_reindex: bool = False
    ) -> BatchFileImportReport:
        target_dir = Path(directory_path)
        if not target_dir.exists():
            return BatchFileImportReport(
                success=False,
                total_files_scanned=0,
                files_processed=0,
                total_articles_indexed=0,
                results=[],
                message=f"Papka topilmadi: {directory_path}"
            )

        supported_extensions = {".docx", ".doc", ".pdf", ".txt", ".md"}
        all_files = [p for p in target_dir.glob("*") if p.is_file() and p.suffix.lower() in supported_extensions]

        results: List[FileImportResult] = []
        total_indexed = 0

        for file_path in sorted(all_files):
            # Skip massive raw HTML table export files unless they contain full law text
            # We check if it is a pure catalog list vs full-text document
            try:
                articles, law_number, law_title = self.parse_file(file_path)
                if articles:
                    await self.store.index_articles_batch(articles)
                    total_indexed += len(articles)
                    results.append(FileImportResult(
                        filename=file_path.name,
                        file_type=file_path.suffix.lower(),
                        articles_count=len(articles),
                        law_number=law_number,
                        law_title=law_title,
                        status="success"
                    ))
                else:
                    results.append(FileImportResult(
                        filename=file_path.name,
                        file_type=file_path.suffix.lower(),
                        articles_count=0,
                        law_number="N/A",
                        law_title="Hujjatlar ro'yxati (katalog) yoki moddalar topilmadi",
                        status="skipped"
                    ))
            except Exception as e:
                logger.error(f"Faylni qayta ishlashda xatolik {file_path.name}: {e}", exc_info=True)
                results.append(FileImportResult(
                    filename=file_path.name,
                    file_type=file_path.suffix.lower(),
                    articles_count=0,
                    law_number="ERROR",
                    law_title="Xatolik",
                    status="error",
                    error=str(e)
                ))

        processed_count = len([r for r in results if r.status == "success"])
        return BatchFileImportReport(
            success=True,
            total_files_scanned=len(all_files),
            files_processed=processed_count,
            total_articles_indexed=total_indexed,
            results=results,
            message=f"{len(all_files)} ta fayl ko'rib chiqildi, {processed_count} ta hujjatdan jami {total_indexed} ta norma indekslandi."
        )

    def parse_file(self, file_path: Path) -> Tuple[List[CanonicalArticle], str, str]:
        ext = file_path.suffix.lower()
        if ext == ".docx":
            return self._parse_docx(file_path)
        elif ext == ".pdf":
            return self._parse_pdf(file_path)
        elif ext == ".doc":
            return self._parse_doc(file_path)
        elif ext in [".txt", ".md"]:
            return self._parse_text_file(file_path)
        return [], "", ""

    def _parse_docx(self, file_path: Path) -> Tuple[List[CanonicalArticle], str, str]:
        import docx
        doc = docx.Document(str(file_path))
        lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                lines.append(text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    lines.append(row_text)

        full_raw_text = "\n".join(lines)
        return self._segment_law_text(full_raw_text, file_path.stem, str(file_path))

    def _parse_pdf(self, file_path: Path) -> Tuple[List[CanonicalArticle], str, str]:
        from pypdf import PdfReader
        reader = PdfReader(str(file_path))
        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pages_text.append(t)
        full_raw_text = "\n".join(pages_text)
        return self._segment_law_text(full_raw_text, file_path.stem, str(file_path))

    def _parse_text_file(self, file_path: Path) -> Tuple[List[CanonicalArticle], str, str]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return self._segment_law_text(text, file_path.stem, str(file_path))

    def _parse_doc(self, file_path: Path) -> Tuple[List[CanonicalArticle], str, str]:
        # In Uzbekistan and Lex.uz, most .doc files are HTML documents or CFBF binary
        try:
            with open(file_path, "rb") as f:
                head = f.read(512)
            if b"<html" in head.lower() or b"<?xml" in head.lower() or b"<!doctype" in head.lower():
                with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                    html_content = f.read()
                soup = BeautifulSoup(html_content, "html.parser")

                # If this is a Lex.uz search results table list, don't parse as single law
                if soup.find("table") and ("Ҳужжат реквизитлари" in soup.get_text() or "Ҳужжат ҳаволаси" in soup.get_text()):
                    logger.info(f"{file_path.name} is a catalog table, skipping statutory segmentation.")
                    return [], "", ""

                for el in soup.select(".lx_elem2, .lx_elem3, .docNabvar, script, style"):
                    el.decompose()
                doc_text = soup.get_text(separator="\n")
                return self._segment_law_text(doc_text, file_path.stem, str(file_path))
            else:
                # Binary DOC fallback
                with open(file_path, "rb") as f:
                    raw_bytes = f.read()
                text_ascii = re.sub(rb"[^\x20-\x7E\r\n\t]", rb" ", raw_bytes).decode("ascii", errors="ignore")
                return self._segment_law_text(text_ascii, file_path.stem, str(file_path))
        except Exception as e:
            logger.warning(f"Error parsing .doc file {file_path}: {e}")
            return [], "", ""

    def _segment_law_text(
        self,
        text: str,
        filename_base: str,
        source_location: str
    ) -> Tuple[List[CanonicalArticle], str, str]:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if not lines:
            return [], "", ""

        # Extract Law Number
        law_number = "NORMA-LOCAL"
        law_title = filename_base.replace("_", " ")

        for line in lines[:15]:
            m = re.search(r"((?:O[‘'`ʼʼ]?RQ|PQ|ПҚ|VMQ|ВМҚ|PF|ПФ|Farmo[ny]ish)[-–—\s]*\d+(?:-[IVX]+)?)", line, re.IGNORECASE)
            if m:
                law_number = m.group(1).replace(" ", "-")
                break

        for line in lines[:15]:
            if len(line) > 15 and not line.lower().startswith("o‘zbekiston respublikasi") and not line.startswith("["):
                law_title = line
                break

        is_cyrillic = detect_script(text[:3000]) == "cyrillic"
        law_title_uz = cyrillic_to_latin(law_title) if is_cyrillic else normalize_apostrophes(law_title)
        law_slug = re.sub(r"[^a-z0-9_]", "_", filename_base.lower()).strip("_")

        # Category and disability detection
        category = self._detect_category(text[:10000])
        disability_types = self._detect_disability_types(text[:10000])

        article_header_regex = re.compile(
            r"^(?:(\d+(?:[-–—\.]\d+|[¹²³⁴⁵⁶⁷⁸⁹⁰]+)?)\s*[-–—]?\s*(?:modda|статья|band|банд|пункт)[\.\s:]+)(.*)$",
            re.IGNORECASE
        )

        articles: List[CanonicalArticle] = []
        current_art_num: Optional[str] = None
        current_art_title: str = ""
        current_lines: List[str] = []

        def flush_article():
            nonlocal current_art_num, current_art_title, current_lines
            if not current_art_num or not current_lines:
                return

            body_str = "\n".join(current_lines).strip()
            if len(body_str) < 15:
                current_art_num = None
                current_lines = []
                return

            b_script = detect_script(body_str)
            body_uz = cyrillic_to_latin(body_str) if b_script == "cyrillic" else normalize_apostrophes(body_str)
            title_uz = cyrillic_to_latin(current_art_title) if b_script == "cyrillic" else normalize_apostrophes(current_art_title)

            clean_num = re.sub(r"[^0-9a-zA-Z_\-]", "", current_art_num)
            canonical_id = f"{law_slug}_art_{clean_num}"

            # Points
            pt_regex = re.compile(r"^\s*(\d+)[\.\)]\s+(.*)$")
            points: List[ArticlePoint] = []
            for l in current_lines:
                pm = pt_regex.match(l)
                if pm:
                    pid = pm.group(1)
                    pt_text = pm.group(2)
                    pt_uz = cyrillic_to_latin(pt_text) if b_script == "cyrillic" else normalize_apostrophes(pt_text)
                    points.append(ArticlePoint(point_id=pid, text=localized(pt_uz, pt_text)))

            if not points:
                points.append(ArticlePoint(point_id="1", text=localized(body_uz, body_str)))

            hasher = hashlib.sha256()
            hasher.update(body_uz.encode("utf-8"))
            s_hash = f"sha256:{hasher.hexdigest()[:16]}"
            today = date.today().isoformat()

            art = CanonicalArticle(
                canonical_id=canonical_id,
                schema_version="1.1",
                metadata=ArticleMetadata(
                    law_id=law_slug,
                    law_title=localized(law_title_uz, law_title),
                    law_number=law_number,
                    chapter_number="",
                    article_number=f"{current_art_num}-modda",
                    article_title=localized(title_uz or f"{current_art_num}-modda", current_art_title),
                    source_url=f"local://{Path(source_location).name}",
                    effective_date=today,
                    last_verified_date=today,
                    source_hash=s_hash,
                    status=ArticleStatus.ACTIVE,
                    category=category,
                    disability_types=disability_types,
                    beneficiary_type=["person_with_disability", "individual"],
                    cross_references=[],
                    embedding_model_version="baai/bge-m3",
                ),
                points=points,
                translations=localized(body_uz, body_str),
                translation_provenance=provenance_for(localized(body_uz, body_str))
            )
            articles.append(art)
            current_art_num = None
            current_art_title = ""
            current_lines = []

        for line in lines:
            m = article_header_regex.match(line)
            if m:
                flush_article()
                current_art_num = m.group(1)
                current_art_title = m.group(2).strip()
            elif current_art_num:
                current_lines.append(line)

        flush_article()
        return articles, law_number, law_title_uz

    def _detect_category(self, text: str) -> str:
        text_low = text.lower()
        if any(w in text_low for w in ["tmek", "тмек", "tiek", "тиек", "tibbiy-ijtimoiy", "тиббий-ижтимоий"]):
            return "medical_social_expertise"
        if any(w in text_low for w in ["pensiya", "пенсия", "nafaqa", "нафақа", "kompensatsiya", "компенсация"]):
            return "pensions_and_benefits"
        if any(w in text_low for w in ["mehnat", "меҳнат", "ish o'rni", "bandlik", "бандлик", "kvota", "квота"]):
            return "labor_rights"
        if any(w in text_low for w in ["soliq", "солиқ", "imtiyoz", "имтиёз"]):
            return "tax_benefits"
        if any(w in text_low for w in ["protez", "протез", "reabilitatsiya", "реабилитация", "aravacha", "аравача"]):
            return "rehabilitation_prosthetics"
        if any(w in text_low for w in ["sanatoriy", "санаторий", "davolanish", "йўлланма", "yo'llanma"]):
            return "sanatorium_health"
        if any(w in text_low for w in ["inklyuziv", "инклюзив", "ta'lim", "таълим", "maktab", "мактаб"]):
            return "inclusive_education"
        if any(w in text_low for w in ["to'siqsiz", "тўсиқсиз", "pandus", "пандус", "shaharsozlik", "шаҳарсозлик", "infratuzilma"]):
            return "accessibility_infrastructure"
        return "disability_rights"

    def _detect_disability_types(self, text: str) -> List[str]:
        types = set()
        text_low = text.lower()
        if "1-guruh" in text_low or "i guruh" in text_low or "i гуруҳ" in text_low:
            types.add("group_1")
        if "2-guruh" in text_low or "ii guruh" in text_low or "ii гуруҳ" in text_low:
            types.add("group_2")
        if "3-guruh" in text_low or "iii guruh" in text_low or "iii гуруҳ" in text_low:
            types.add("group_3")
        if "bola" in text_low or "бола" in text_low:
            types.add("childhood_disability")
        if "ko'rish" in text_low or "кўриш" in text_low or "ko'zi ojiz" in text_low:
            types.add("visual")
        if "eshitish" in text_low or "эшитиш" in text_low or "kar" in text_low:
            types.add("hearing")
        if "harakat" in text_low or "ҳаракат" in text_low or "tayanch" in text_low:
            types.add("physical")
        if not types:
            types.add("general")
        return sorted(list(types))
